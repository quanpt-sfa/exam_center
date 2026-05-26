"""Parsers for CSV/XLSX import files."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.errors import ApiError


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower()


def _normalize_cell(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    return value


class FileParserService:
    """Parse supported import files into normalized row dictionaries."""

    def parse(self, file_path: str, original_filename: str) -> dict:
        path = Path(file_path)
        if not path.exists():
            raise ApiError(status_code=404, code="import_file_not_found", message="Upload file not found", details={})

        suffix = path.suffix.lower()
        if suffix == ".csv":
            return self._parse_csv(path)
        if suffix in {".xlsx", ".xlsm"}:
            return self._parse_xlsx(path)

        raise ApiError(
            status_code=400,
            code="unsupported_import_file_type",
            message="Only CSV and XLSX files are supported",
            details={"filename": original_filename},
        )

    def _parse_csv(self, path: Path) -> dict:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames is None:
                raise ApiError(
                    status_code=400,
                    code="import_parse_error",
                    message="CSV file must include a header row",
                    details={"file": path.name},
                )

            headers = [_normalize_header(name) for name in reader.fieldnames]
            rows: list[dict] = []
            for row in reader:
                normalized = {_normalize_header(k): _normalize_cell(v) for k, v in row.items()}
                rows.append(normalized)

        return {"sheet_name": "csv", "headers": headers, "rows": rows}

    def _parse_xlsx(self, path: Path) -> dict:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]

        iterator = worksheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if not header_row:
            raise ApiError(
                status_code=400,
                code="import_parse_error",
                message="XLSX file must include a header row",
                details={"file": path.name},
            )

        headers = [_normalize_header(item) for item in header_row]
        rows: list[dict] = []
        for values in iterator:
            row = {}
            for index, header in enumerate(headers):
                if not header:
                    continue
                cell = values[index] if index < len(values) else None
                row[header] = _normalize_cell(cell)
            rows.append(row)

        return {"sheet_name": worksheet.title, "headers": headers, "rows": rows}
